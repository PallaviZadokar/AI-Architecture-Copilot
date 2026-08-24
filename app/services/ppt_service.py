from io import BytesIO

from pptx import Presentation
from pptx.util import Inches, Pt


def add_title(
    slide,
    title: str,
    subtitle: str = "",
):

    title_box = slide.shapes.add_textbox(
        Inches(0.5),
        Inches(0.3),
        Inches(12),
        Inches(0.7),
    )

    title_frame = title_box.text_frame

    paragraph = title_frame.paragraphs[0]

    paragraph.text = title

    paragraph.font.size = Pt(26)

    paragraph.font.bold = True

    if subtitle:

        subtitle_box = slide.shapes.add_textbox(
            Inches(0.5),
            Inches(1.0),
            Inches(12),
            Inches(0.4),
        )

        subtitle_box.text_frame.text = subtitle


def create_architecture_ppt(
    result: dict,
) -> bytes:

    prs = Presentation()

    architecture = result.get(
        "architecture",
        {},
    )

    technology = result.get(
        "technology_stack",
        {},
    )

    validation = result.get(
        "validation",
        {},
    )

    # Slide 1
    slide = prs.slides.add_slide(
        prs.slide_layouts[6]
    )

    add_title(
        slide,
        "Architecture Copilot",
        architecture.get(
            "project_summary",
            "Solution Architecture",
        ),
    )

    textbox = slide.shapes.add_textbox(
        Inches(0.8),
        Inches(2),
        Inches(11),
        Inches(2),
    )

    textbox.text_frame.text = (
        "High-Level Solution Architecture"
    )

    # Slide 2
    slide = prs.slides.add_slide(
        prs.slide_layouts[6]
    )

    add_title(
        slide,
        "Architecture Overview",
    )

    y = 1.4

    for component in architecture.get(
        "components",
        [],
    ):

        box = slide.shapes.add_textbox(
            Inches(0.8),
            Inches(y),
            Inches(11),
            Inches(0.65),
        )

        text = (
            f"{component.get('name')} | "
            f"{component.get('technology')}"
        )

        box.text_frame.text = text

        y += 0.75

        if y > 6.8:
            break

    # Slide 3
    slide = prs.slides.add_slide(
        prs.slide_layouts[6]
    )

    add_title(
        slide,
        "Technology Decisions",
    )

    y = 1.4

    for item in technology.get(
        "recommendations",
        [],
    ):

        box = slide.shapes.add_textbox(
            Inches(0.7),
            Inches(y),
            Inches(12),
            Inches(0.75),
        )

        text = (
            f"{item.get('category')}: "
            f"{item.get('technology')}\n"
            f"{item.get('description')}"
        )

        box.text_frame.text = text

        y += 0.9

        if y > 6.5:
            break

    # Slide 4
    slide = prs.slides.add_slide(
        prs.slide_layouts[6]
    )

    add_title(
        slide,
        "Architecture Validation",
    )

    textbox = slide.shapes.add_textbox(
        Inches(0.7),
        Inches(1.4),
        Inches(11),
        Inches(4.8),
    )

    frame = textbox.text_frame

    frame.text = (
        f"Status: "
        f"{validation.get('overall_status', 'N/A')}"
    )

    for section in [
        "strengths",
        "issues",
        "risks",
        "recommendations",
    ]:

        paragraph = frame.add_paragraph()

        paragraph.text = (
            "\n"
            + section.upper()
        )

        paragraph.font.bold = True

        for item in validation.get(
            section,
            [],
        ):

            paragraph = frame.add_paragraph()

            paragraph.text = (
                f"• {item}"
            )

    output = BytesIO()

    prs.save(output)

    return output.getvalue()
